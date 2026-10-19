from pathlib import Path

from docx import Document

from recipify import pipeline
from recipify.categories import Category
from recipify.classify import classify
from recipify.metadata import Metadata
from recipify.naming import rename_plan, slugify, standard_name
from recipify.quality import REVIEW_THRESHOLD, Quality, detect_language, score_text


def test_slugify_icelandic():
    assert slugify("Þorskur & hörpuskel (tasting)") == "thorskur-horpuskel-tasting"


def test_standard_name_copy_and_suffix():
    assert standard_name(Path("Advanced Pastry 10 (1).JPEG")) == "advanced-pastry-10-copy1.jpg"


def test_rename_plan_resolves_collisions(tmp_path: Path):
    a, b = tmp_path / "A b.pdf", tmp_path / "a-b.pdf"
    new_names = {
        p.name if p not in dict(rename_plan([a, b])) else dict(rename_plan([a, b]))[p]
        for p in (a, b)
    }
    assert len(new_names) == 2


def make_docx(path: Path, text: str) -> None:
    doc = Document()
    if text:
        doc.add_paragraph(text)
    doc.save(str(path))


def test_classify(tmp_path: Path):
    (tmp_path / "empty.pdf").write_bytes(b"")
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    (tmp_path / "~$lock.docx").write_bytes(b"x")
    (tmp_path / "a.odt").write_bytes(b"x")
    (tmp_path / "p.jpg").write_bytes(b"\xff\xd8\xff\xe0rest")
    make_docx(tmp_path / "ok.docx", "Pancakes")
    make_docx(tmp_path / "blank.docx", "")
    expect = {
        "empty.pdf": Category.EMPTY_OR_CORRUPTED,
        "bad.pdf": Category.EMPTY_OR_CORRUPTED,
        "~$lock.docx": Category.IGNORED,
        "a.odt": Category.LEGACY_OFFICE,
        "p.jpg": Category.IMAGE_FILES,
        "ok.docx": Category.DOCX_FILES,
        "blank.docx": Category.EMPTY_OR_CORRUPTED,
    }
    for name, cat in expect.items():
        assert classify(tmp_path / name).category == cat, name


def test_process_versions_outputs(tmp_path: Path):
    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    make_docx(recipes / "Pönnukökur.docx", "Pancakes")
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)

    pipeline.process(meta, recipes, outputs, run_id="run1")
    assert (outputs / "run1" / "docx_files" / "ponnukokur.md").read_text().strip() == "Pancakes"
    pipeline.process(meta, recipes, outputs, run_id="run2")  # unchanged: skipped
    pipeline.process(meta, recipes, outputs, force=True, run_id="run3")
    latest = Metadata.load_latest(outputs)
    versions = latest.recipes["Pönnukökur.docx"]["versions"]
    assert [v["run_id"] for v in versions] == ["run1", "run3"]
    assert latest.recipes["Pönnukökur.docx"]["original_format"] == "docx"
    # every run keeps its own metadata snapshot and changelog
    assert (
        len(Metadata(outputs / "run1" / "metadata.json").recipes["Pönnukökur.docx"]["versions"])
        == 1
    )
    assert "Pönnukökur.docx" in (outputs / "run3" / "CHANGELOG.md").read_text()
    assert not (outputs / "run2").exists()
    assert (outputs / "run1").exists() and (outputs / "run3").exists()


def test_register_manual_adds_version(tmp_path: Path):
    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    make_docx(recipes / "a.docx", "x")
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    md = tmp_path / "m.md"
    md.write_text("# hi\n")
    q = Quality(0.9, method="manual", notes="checked")
    out = pipeline.register_manual(
        meta, outputs, "a.docx", md, "copilot-vision", run_id="r1", language="is", quality=q
    )
    assert (outputs / out).read_text() == "# hi\n"
    v = Metadata(outputs / "r1" / "metadata.json").recipes["a.docx"]["versions"][0]
    assert v["processor"] == "copilot-vision"
    assert v["quality"] == {"score": 0.9, "method": "manual", "flags": [], "notes": "checked"}
    assert (outputs / "r1" / "CHANGELOG.md").exists()


ICELANDIC = (
    "Látið malla í tuttugu mínútur. Setjið smjör og salt í pottinn með hveiti og sykri."
)


def test_detect_language():
    assert detect_language(ICELANDIC) == "is"
    assert (
        detect_language("Mix the flour and sugar with the butter until smooth, then bake.") == "en"
    )
    assert detect_language("hi") == "unknown"


def test_quality_scores_garbage_low():
    good = score_text(ICELANDIC * 3)
    garbled = score_text("xkqzt wrbnp ghjkl p q r s t u v w x y z mnbvc fdsgh " * 4)
    spaced = score_text("S o ð i ð n i ð u r þ a r t i l b r a g ð i ð e r " * 3)
    mojibake = score_text("Sm\ufffdr og sl\ufffdt \ufffd pottinn með salti og pipar og hveiti " * 3)
    assert good.score >= 0.9 and not good.flags
    assert garbled.score < REVIEW_THRESHOLD and "nonwords" in garbled.flags
    assert spaced.score < REVIEW_THRESHOLD and "broken_spacing" in spaced.flags
    assert "mojibake" in mojibake.flags
    assert score_text("").score < REVIEW_THRESHOLD


def test_new_version_carries_forward_previous_outputs(tmp_path: Path):
    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    make_docx(recipes / "a.docx", "Alpha")
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    pipeline.process(meta, recipes, outputs, run_id="run1")

    make_docx(recipes / "b.docx", "Beta")
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    pipeline.process(meta, recipes, outputs, run_id="run2")

    # run2 is a complete tree: the unchanged file is copied, the new one converted
    assert (outputs / "run2" / "docx_files" / "a.md").read_text().strip() == "Alpha"
    assert (outputs / "run2" / "docx_files" / "b.md").read_text().strip() == "Beta"
    assert (outputs / "run1" / "docx_files" / "a.md").exists()  # history untouched
    latest = Metadata.load_latest(outputs)
    carried = latest.recipes["a.docx"]["versions"][-1]
    assert carried["carried_from"] == "run1" and carried["output"].startswith("run2/")
    log = (outputs / "run2" / "CHANGELOG.md").read_text()
    assert "1 new recipe file(s) converted" in log and "1 unchanged output(s)" in log


def test_spreadsheet_to_markdown(tmp_path: Path):
    from openpyxl import Workbook

    from recipify.extractors.office import spreadsheet_to_markdown

    wb = Workbook()
    ws = wb.active
    ws.title = "Dough"
    ws.append(["Ingredient", "Grams"])
    ws.append(["Flour", 500.0])
    ws.append([None, None])
    wb.save(tmp_path / "dough.xlsx")
    md = spreadsheet_to_markdown(tmp_path / "dough.xlsx")
    assert "## Dough" in md and "| Flour | 500 |" in md


def test_misnamed_zip_formats_are_sniffed(tmp_path: Path):
    import zipfile

    from recipify.extractors.epub import epub_to_markdown

    epub = tmp_path / "book.pdf"  # an EPUB with a .pdf extension
    with zipfile.ZipFile(epub, "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
        zf.writestr(
            "META-INF/container.xml",
            '<container><rootfiles><rootfile full-path="OEBPS/c.opf"/></rootfiles></container>',
        )
        zf.writestr(
            "OEBPS/c.opf",
            '<package xmlns="http://www.idpf.org/2007/opf"><manifest>'
            '<item id="a" href="a.html"/></manifest><spine><itemref idref="a"/></spine></package>',
        )
        zf.writestr("OEBPS/a.html", "<html><body><h1>Soup</h1><p>Boil water.</p></body></html>")
    assert classify(epub).category == Category.EPUB_BOOKS
    md = epub_to_markdown(epub)
    assert "# Soup" in md and "Boil water." in md
    make_docx(tmp_path / "menu.pdf", "Menu text")
    assert classify(tmp_path / "menu.pdf").category == Category.DOCX_FILES


def test_viewer_builds_versioned_single_file(tmp_path: Path):
    from recipify import viewer

    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    for i in range(3):
        make_docx(recipes / f"r{i}.docx", f"Recipe {i} with </script> flour " * 20 + str(i))
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    pipeline.process(meta, recipes, outputs, run_id="run1")

    pipeline.enrich(meta, outputs, {"titles": {}, "labels": {}, "kind": {}}, run_id="run2")
    first = viewer.build(outputs, tmp_path / "viewer", per_group=2)
    second = viewer.build(outputs, tmp_path / "viewer", per_group=2)
    assert first.parent.name == "v001" and second.parent.name == "v002"
    page = first.read_text(encoding="utf-8")
    assert page.count("</script>") == 2  # data block + app script: payload cannot break out
    assert "__DATA__" not in page and "http" not in page.split("<script")[0]
    assert "2 sample recipes" in (first.parent / "NOTES.md").read_text()


def test_title_refinement_and_labels():
    from recipify.labels import detect_labels, fix_mojibake, refine_title

    assert fix_mojibake("BBQ S¢sa") == "BBQ Sósa"
    assert fix_mojibake("Eftirr\x82tta b¢k") == "Eftirrétta bók"
    assert fix_mojibake("Sósa") == "Sósa"  # already fine
    assert (
        refine_title("Untitled presentation-2", "# x\n\nA la carte æfing\n", "u", {})
        == "A la carte æfing"
    )
    assert refine_title("CHICHARRON EMPANADA DOUGH", "", "c", {}) == "Chicharron empanada dough"
    assert (
        refine_title(
            "20200101120000000", "", "20200101120000000", {"20200101120000000": "Notebook"}
        )
        == "Notebook"
    )
    text = "Hrísgrjón\n200 g hrísgrjón\n1 laukur\nsveppir 100 g\nsalt og pipar, vatn, olía " * 3
    lab = detect_labels("Sveppa risotto", "Eldhús/Aðalréttir/x.docx", text)
    assert (
        "main" in lab["course"] and "rice" in lab["dish"] and lab["diet"] == ["vegetarian", "vegan"]
    )
    meat = detect_labels("Lambasteik", "x.docx", "lambakjöt 500 g, salt, pipar, smjör " * 10)
    assert "lamb" in meat["ingredients"] and meat["diet"] == ["meat"]


def test_enrich_adds_labels_and_keeps_original_name(tmp_path: Path):
    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    make_docx(
        recipes / "Untitled.docx", "# Untitled\n\nSúkkulaðikaka\n\n200 g sykur, 100 g smjör, kaka"
    )
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    pipeline.process(meta, recipes, outputs, run_id="run1")
    assert pipeline.enrich(meta, outputs, {}, run_id="run2") == "run2"
    info = Metadata.load_latest(outputs).recipes["Untitled.docx"]["recipe"]
    assert info["original_filename"] == "Untitled.docx" and info["title"] != "Untitled"
    new_file = outputs / "run2" / "docx_files" / "untitled.md"
    assert new_file.read_text().startswith(f"# {info['title']}")
    assert (outputs / "run1" / "docx_files" / "untitled.md").read_text().startswith("# Untitled")
    assert pipeline.enrich(Metadata.load_latest(outputs), outputs, {}, run_id="run3") is None


def test_tidy_removes_echoes_and_repeats_but_keeps_raw_version(tmp_path: Path):
    from recipify.tidy import tidy

    raw = "# K b\n\n# K b\n\n## K b\n\n- Salat\n\n300 gr. Sýrður rjómi\n"
    assert tidy(raw, "K b") == "# K b\n\n- Salat\n\n300 gr. Sýrður rjómi\n"
    cells = (
        "Lunch Menu 1 | Hádegi Lunch Menu 1 | Hádegi Lunch Menu 1 | Hádegi Lunch Menu 1 | Hádegi"
    )
    assert tidy(cells, "X").endswith("Lunch Menu 1 | Hádegi\n")
    page = "\n\n".join(
        ["Línan eitt hér", "Önnur lína með texta í", "Þriðja línan er löng og skýr"] * 4
    )
    assert tidy(page, "X").count("Línan eitt hér") == 1

    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    make_docx(recipes / "a.docx", "Soup\n\nSoup\n\n200 g onion")
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    pipeline.process(meta, recipes, outputs, run_id="run1")
    pipeline.enrich(meta, outputs, {}, run_id="run2")
    assert pipeline.tidy_run(meta, outputs, run_id="run3") in ("run3", None)
    assert pipeline.tidy_run(Metadata.load_latest(outputs), outputs, run_id="run4") is None


def test_spelling_fixes_keep_case_and_raw_version(tmp_path: Path):
    from recipify.spelling import fix_spelling

    fixes = {"sírtónu": "sítrónu", "smjór": "smjör"}
    assert fix_spelling("Smjór og SÍRTÓNU, sírtónu.", fixes) == ("Smjör og SÍTRÓNU, sítrónu.", 3)
    assert fix_spelling("Smjörið", fixes)[1] == 0  # whole words only

    recipes, outputs = tmp_path / "recipes", tmp_path / "outputs"
    recipes.mkdir()
    make_docx(recipes / "a.docx", "Kaka\n\n200 g smjór og sírtónu")
    meta = Metadata.load_latest(outputs)
    pipeline.classify_all(meta, recipes)
    pipeline.process(meta, recipes, outputs, run_id="run1")
    assert pipeline.spelling_run(meta, outputs, fixes, run_id="run2") == "run2"
    assert "smjör og sítrónu" in (outputs / "run2" / "docx_files" / "a.md").read_text()
    assert "smjór" in (outputs / "run1" / "docx_files" / "a.md").read_text()
    assert (
        pipeline.spelling_run(Metadata.load_latest(outputs), outputs, fixes, run_id="run3") is None
    )


def test_locked_page_roundtrip():
    import pytest
    from cryptography.exceptions import InvalidTag

    from recipify.lock import decrypt_page, encrypt_page

    page = "<!DOCTYPE html><p>Kjötsúpa</p>"
    locked = encrypt_page(page, "s3cret/+&=key")
    assert "Kjötsúpa" not in locked and locked.count("</script>") == 2
    assert decrypt_page(locked, "s3cret/+&=key") == page
    with pytest.raises(InvalidTag):
        decrypt_page(locked, "wrong")
