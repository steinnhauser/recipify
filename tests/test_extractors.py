import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from recipify.extractors import ocr, office
from recipify.extractors.docx import docx_to_markdown
from recipify.extractors.epub import _html_to_text
from recipify.extractors.pdf import pdf_to_markdown


def test_docx_styles_and_tables(tmp_path: Path):
    doc = Document()
    doc.add_heading("Kjötsúpa", level=0)  # the "Title" style
    doc.add_heading("Ingredients", level=2)
    doc.add_paragraph("Lamb", style="List Bullet")
    doc.add_paragraph("")
    doc.add_paragraph("Simmer for an hour.")
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(
        table._cells, ["Item", "Grams", "Rófur | gulrætur", "300\ng"], strict=True
    ):
        cell.text = text
    doc.add_table(rows=0, cols=1)  # empty tables are dropped
    doc.save(str(tmp_path / "soup.docx"))

    assert docx_to_markdown(tmp_path / "soup.docx") == (
        "# Kjötsúpa\n\n## Ingredients\n\n- Lamb\n\nSimmer for an hour.\n\n"
        "| Item | Grams |\n| --- | --- |\n| Rófur \\| gulrætur | 300 g |\n"
    )


def test_pdf_pages_become_sections(tmp_path: Path, make_pdf):
    make_pdf(tmp_path / "one.pdf", ["Only page"])
    two = make_pdf(tmp_path / "two.pdf", ["First", "Second"])
    locked = PdfWriter(clone_from=two)
    locked.encrypt("", owner_password="owner")
    locked.write(tmp_path / "locked.pdf")

    assert pdf_to_markdown(tmp_path / "one.pdf") == "# one\n\nOnly page\n"
    expected = "## Page 1\n\nFirst\n\n## Page 2\n\nSecond\n"
    assert pdf_to_markdown(two) == "# two\n\n" + expected
    assert pdf_to_markdown(tmp_path / "locked.pdf") == "# locked\n\n" + expected


def test_epub_html_to_text():
    page = (
        "<html><head><title>skip</title><style>p{}</style></head><body>"
        "<h2>Sauce</h2><script>var x</script><ul><li>Butter</li><li>Salt</li></ul>"
        "<p>Melt   slowly.</p></body></html>"
    )
    assert _html_to_text(page) == "## Sauce\n\n- Butter\n\n- Salt\n\nMelt slowly."


def test_spreadsheet_cells(tmp_path: Path):
    wb = Workbook()
    wb.active.append(["Sugar", 1.23456789, None, "a\nb"])
    wb.create_sheet("Empty")
    wb.save(tmp_path / "s.xlsx")
    md = office.spreadsheet_to_markdown(tmp_path / "s.xlsx")
    assert "| Sugar | 1.2346 |  | a b |" in md and "## Empty" not in md


@pytest.fixture
def soffice(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    """Fake LibreOffice: 'converts' by copying a prepared file of the requested format."""
    made: dict[str, Path] = {}
    calls: list[list[str]] = []

    def run(cmd: list[str], **kwargs: object) -> None:
        calls.append(cmd)
        target, outdir, src = cmd[cmd.index("--convert-to") + 1], Path(cmd[-2]), Path(cmd[-1])
        if target in made:
            shutil.copyfile(made[target], outdir / f"{src.stem}.{target}")

    monkeypatch.setattr(office.shutil, "which", lambda name: "/usr/bin/" + name)
    monkeypatch.setattr(office.subprocess, "run", run)
    return SimpleNamespace(made=made, calls=calls)


def test_legacy_formats_via_libreoffice(tmp_path: Path, soffice, make_docx, make_pdf):
    wb = Workbook()
    wb.active.append(["Flour", 500])
    wb.save(tmp_path / "made.xlsx")
    soffice.made |= {
        "xlsx": tmp_path / "made.xlsx",
        "docx": make_docx(tmp_path / "made.docx", "Old recipe"),
        "pdf": make_pdf(tmp_path / "made.pdf", ["Slide text"]),
    }
    for name in ("costs.xls", "old.doc", "talk.ppt", "titled.odt"):
        (tmp_path / name).write_bytes(b"legacy")

    assert "| Flour | 500 |" in office.spreadsheet_to_markdown(tmp_path / "costs.xls")
    assert office.legacy_office_to_markdown(tmp_path / "old.doc") == "# old\n\nOld recipe\n"
    assert "Slide text" in office.legacy_office_to_markdown(tmp_path / "talk.ppt")
    doc = Document()
    doc.add_heading("Already titled", level=1)
    doc.save(str(soffice.made["docx"]))
    assert office.legacy_office_to_markdown(tmp_path / "titled.odt") == "# Already titled\n"
    assert all("--headless" in c for c in soffice.calls)


def test_libreoffice_errors(tmp_path: Path, soffice, monkeypatch: pytest.MonkeyPatch):
    (tmp_path / "old.doc").write_bytes(b"legacy")
    with pytest.raises(RuntimeError, match="produced no docx"):
        office.legacy_office_to_markdown(tmp_path / "old.doc")
    monkeypatch.setattr(office.shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="LibreOffice"):
        office.legacy_office_to_markdown(tmp_path / "old.doc")


class FakeMistral:
    sent: list[dict] = []

    def __init__(self, api_key: str) -> None:
        self.ocr = self

    def process(self, document: dict, model: str, include_image_base64: bool) -> SimpleNamespace:
        FakeMistral.sent.append(document)
        return SimpleNamespace(
            pages=[SimpleNamespace(markdown="# Scan"), SimpleNamespace(markdown="Text ")]
        )


def test_ocr(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(ocr, "load_dotenv", lambda: None)  # never read the developer's .env
    monkeypatch.delenv("MISTRAL_API_KEY", raising=False)
    monkeypatch.delenv("API_KEY", raising=False)
    (tmp_path / "scan.pdf").write_bytes(b"%PDF")
    (tmp_path / "photo.JPG").write_bytes(b"\xff\xd8\xff")
    with pytest.raises(RuntimeError, match="MISTRAL_API_KEY"):
        ocr.ocr_to_markdown(tmp_path / "scan.pdf")

    monkeypatch.setenv("MISTRAL_API_KEY", "k")
    monkeypatch.setattr(ocr, "Mistral", FakeMistral)
    assert ocr.ocr_to_markdown(tmp_path / "scan.pdf") == "# Scan\n\nText\n"
    ocr.ocr_to_markdown(tmp_path / "photo.JPG")
    pdf, photo = FakeMistral.sent[-2:]
    assert pdf["type"] == "document_url" and pdf["document_url"].startswith("data:application/pdf")
    assert photo["type"] == "image_url" and photo["image_url"].startswith("data:image/jpeg")
