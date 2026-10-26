from collections.abc import Callable
from pathlib import Path

import pytest
from docx import Document

from recipify.store import Workspace

EXAMPLES = Path(__file__).parent.parent / "examples" / "sources"


def _make_pdf(path: Path, pages: list[str]) -> Path:
    """A minimal PDF, one line of Helvetica text per line of each page (no PDF library)."""
    n = len(pages)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {n} >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
    ]
    for text in pages:
        lines = "".join(f"0 -16 Td ({line}) Tj " for line in text.splitlines())
        stream = f"BT /F1 12 Tf 72 760 Td {lines}ET"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objects) + 2} 0 R >>"
        )
        objects.append(f"<< /Length {len(stream.encode('cp1252'))} >>\nstream\n{stream}\nendstream")
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("cp1252")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets).encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(out)
    return path


def _make_docx(path: Path, *paragraphs: str) -> Path:
    """Paragraphs; a newline inside one becomes a soft line break, as in kitchen files."""
    doc = Document()
    for text in paragraphs:
        run = doc.add_paragraph().add_run()
        for i, line in enumerate(text.split("\n")):
            if i:
                run.add_break()
            run.add_text(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


@pytest.fixture
def make_pdf() -> Callable[[Path, list[str]], Path]:
    return _make_pdf


@pytest.fixture
def make_docx() -> Callable[..., Path]:
    return _make_docx


TART = """Tartdeig
125gr smjör
2 egg
340gr kökuhveiti
Aðferð:
Allt sett saman í hrærivél og unnið
saman með hnoðara.
Bakað á 160°C í 28 mínútur."""


@pytest.fixture
def ws(tmp_path: Path) -> Workspace:
    """A workspace with a few sources: a recipe PDF, a recipe docx, a copy, a book, a scan."""
    sources = tmp_path / "recipes"
    _make_pdf(sources / "Bakarí" / "Tartdeig.pdf", [TART])
    _make_pdf(sources / "copy" / "Tartdeig (1).pdf", [TART])
    _make_docx(
        sources / "Eftirréttir" / "Skyrkaka.docx",
        "Skyrkaka",
        "500 g skyr\n200 g rjómi\n60 g sykur",
        "Aðferð:",
        "Allt hrært saman og kælt yfir nótt.",
    )
    _make_pdf(sources / "Handbook.pdf", [f"Page {i}" for i in range(10)])
    (sources / "photo.jpg").write_bytes(b"\xff\xd8\xff\xe0 not really a photo")
    return Workspace(sources, tmp_path / "data")
