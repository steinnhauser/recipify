import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from openpyxl import Workbook

from recipify.categories import Category
from recipify.classify import classify
from recipify.extract import LOCAL, office
from recipify.extract.docx import extract_docx
from recipify.extract.pdf import _poppler, extract_pdf, order_agreement, page_count, render_pages


def test_classify(tmp_path: Path, make_pdf, make_docx):
    (tmp_path / "empty.pdf").write_bytes(b"")
    (tmp_path / "bad.pdf").write_bytes(b"not a pdf")
    (tmp_path / "broken.pdf").write_bytes(b"%PDF-1.4\ngarbage")
    (tmp_path / "~$lock.docx").write_bytes(b"x")
    (tmp_path / "a.odt").write_bytes(b"x")
    (tmp_path / "s.xlsx").write_bytes(b"x")
    (tmp_path / "b.epub").write_bytes(b"x")
    (tmp_path / "notes.txt").write_bytes(b"x")
    (tmp_path / "p.jpg").write_bytes(b"\xff\xd8\xff\xe0rest")
    (tmp_path / "fake.png").write_bytes(b"not an image")
    (tmp_path / "zipped.docx").write_bytes(b"PK not really")
    make_docx(tmp_path / "ok.docx", "Pancakes")
    make_docx(tmp_path / "blank.docx")
    make_pdf(tmp_path / "text.pdf", ["\n".join(["A recipe line long enough"] * 5)])
    make_pdf(tmp_path / "scan.pdf", ["x"])
    shutil.copy(tmp_path / "ok.docx", tmp_path / "misnamed.pdf")
    expect = {
        "empty.pdf": Category.EMPTY_OR_CORRUPTED,
        "bad.pdf": Category.EMPTY_OR_CORRUPTED,
        "broken.pdf": Category.EMPTY_OR_CORRUPTED,
        "~$lock.docx": Category.IGNORED,
        "a.odt": Category.LEGACY_OFFICE,
        "s.xlsx": Category.SPREADSHEETS,
        "b.epub": Category.EPUB_BOOKS,
        "notes.txt": Category.OTHER_FORMATS,
        "p.jpg": Category.IMAGE_FILES,
        "fake.png": Category.EMPTY_OR_CORRUPTED,
        "zipped.docx": Category.EMPTY_OR_CORRUPTED,
        "ok.docx": Category.DOCX_FILES,
        "blank.docx": Category.EMPTY_OR_CORRUPTED,
        "text.pdf": Category.READABLE_PDFS,
        "scan.pdf": Category.IMAGE_BASED_PDFS,
        "misnamed.pdf": Category.DOCX_FILES,
    }
    for name, category in expect.items():
        assert classify(tmp_path / name).category == category, name
    assert classify(tmp_path / "text.pdf").pages == 1


def test_classify_epub_and_damaged_docx(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    import zipfile

    with zipfile.ZipFile(tmp_path / "book.pdf", "w") as zf:
        zf.writestr("mimetype", "application/epub+zip")
    with zipfile.ZipFile(tmp_path / "other.pdf", "w") as zf:
        zf.writestr("readme.txt", "hello")
    with zipfile.ZipFile(tmp_path / "damaged.docx", "w") as zf:
        zf.writestr("word/document.xml", "<broken")
    assert classify(tmp_path / "book.pdf").category == Category.EPUB_BOOKS
    assert classify(tmp_path / "other.pdf").category == Category.EMPTY_OR_CORRUPTED
    assert classify(tmp_path / "damaged.docx").category == Category.EMPTY_OR_CORRUPTED

    from recipify import classify as module

    monkeypatch.setattr(module, "page_count", lambda p: 0)
    (tmp_path / "zero.pdf").write_bytes(b"%PDF-1.4")
    monkeypatch.setattr(module, "pdf_text", lambda *a: "")
    assert classify(tmp_path / "zero.pdf").reason == "PDF has no pages"


def test_pdf_extraction_keeps_lines_and_pages(tmp_path: Path, make_pdf):
    pdf = make_pdf(tmp_path / "r.pdf", ["Tartdeig\n125gr smjör", "", "Page two"])
    text, checks = extract_pdf(pdf)
    assert text == "Tartdeig\n125gr smjör\n\f\nPage two\n"
    assert checks == {"order_agreement": 1.0}
    assert page_count(pdf) == 3


def test_order_agreement():
    assert order_agreement("", "") == 1.0
    assert order_agreement("a b\nc", "a b   x\nc") == 1.0
    assert order_agreement("a b\nc d", "a   x b\nc d") == 0.5


def test_poppler_errors(monkeypatch: pytest.MonkeyPatch, tmp_path: Path):
    with pytest.raises(RuntimeError, match="pdfinfo failed"):
        _poppler("pdfinfo", str(tmp_path / "missing.pdf"))
    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="install poppler"):
        _poppler("pdfinfo", "x.pdf")


TEXTBOX = """<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
  xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"
  xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
  <w:r><w:t>Anchor</w:t><w:tab/><w:t>text</w:t></w:r>
  <w:r><mc:AlternateContent>
    <mc:Choice Requires="wps"><wps:txbx><w:txbxContent>
      <w:p><w:r><w:t>Humarsósa</w:t></w:r></w:p>
      <w:p><w:r><w:t>5 stk laukur</w:t><w:br/><w:t>2 l rjómi</w:t></w:r></w:p>
    </w:txbxContent></wps:txbx></mc:Choice>
    <mc:Fallback><w:pict><w:txbxContent><w:p><w:r><w:t>Humarsósa</w:t></w:r></w:p>
    </w:txbxContent></w:pict></mc:Fallback>
  </mc:AlternateContent></w:r>
</w:p>"""
CONTROL = """<w:sdt xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:sdtContent><w:p><w:r><w:t>In a content control</w:t></w:r></w:p></w:sdtContent></w:sdt>"""


def test_docx_extraction(tmp_path: Path):
    doc = Document()
    doc.add_heading("Kjötsúpa", level=0)
    doc.add_heading("Súpan", level=2)
    doc.add_heading("Odd heading", level=1).style.name = "Heading x"
    doc.add_paragraph("Lamb", style="List Bullet")
    doc.add_paragraph("")
    body = doc.element.body
    body.insert(len(body) - 1, parse_xml(TEXTBOX))
    body.insert(len(body) - 1, parse_xml(CONTROL))
    body.insert(
        len(body) - 1,
        parse_xml(CONTROL.replace("<w:sdtContent>", "<w:x>").replace("</w:sdtContent>", "</w:x>")),
    )
    rows = doc.add_table(rows=2, cols=2)
    rows.cell(0, 0).merge(rows.cell(0, 1)).text = "Merged"  # merged cells appear once
    rows.cell(1, 0).text, rows.cell(1, 1).text = "600 g", "cod"
    layout = doc.add_table(rows=1, cols=2)
    layout.cell(0, 0).text = "200 g skyr\n100 g rjómi"
    layout.cell(0, 1).text = "Method"
    doc.add_table(rows=1, cols=1)  # empty tables are dropped
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)  # a page break of its own
    after = doc.add_paragraph("Next page", style="List Bullet")
    after.paragraph_format.page_break_before = True
    doc.save(str(tmp_path / "soup.docx"))
    assert extract_docx(tmp_path / "soup.docx") == (
        "# Kjötsúpa\n\n## Súpan\n\n# Odd heading\n\n- Lamb\n\nAnchor text\n\nHumarsósa\n\n"
        "5 stk laukur\n2 l rjómi\n\nIn a content control\n\nMerged\n600 g | cod\n\n"
        "200 g skyr\n100 g rjómi\n\nMethod\n\n\f\n\n\f\n- Next page\n"
    )


def test_spreadsheets_and_legacy_office(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "Brauð"
    ws.append(["Hveiti", 1000.0, 2.5, None, "a | b"])
    ws.append([None, None])
    wb.create_sheet("Tómt")
    wb.save(tmp_path / "b.xlsx")
    assert office.extract_spreadsheet(tmp_path / "b.xlsx") == (
        "# b\n\n## Brauð\n\nHveiti | 1000 | 2.5 | a / b\n"
    )

    def fake_soffice(path: Path, target: str, outdir: Path) -> Path:
        if target == "xlsx":
            return tmp_path / "b.xlsx"
        if target == "pdf":
            return Path(__file__).parent.parent / "examples/sources/Restaurant/Mains/Lamb shank.pdf"
        out = outdir / "converted.docx"
        Document().save(str(out))
        return out

    assert (office._cell(3.0), office._cell(2.123456), office._cell(None)) == ("3", "2.1235", "")
    monkeypatch.setattr(office, "_soffice", fake_soffice)
    assert office.extract_spreadsheet(tmp_path / "old.xls").startswith("# old\n")
    assert office.extract_legacy(tmp_path / "slides.pptx").startswith("Slow-braised lamb shanks")
    assert office.extract_legacy(tmp_path / "note.doc") == "\n"


def test_soffice_wrapper(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    monkeypatch.setattr(shutil, "which", lambda name: None)
    with pytest.raises(RuntimeError, match="LibreOffice"):
        office._soffice(tmp_path / "a.doc", "docx", tmp_path)
    monkeypatch.setattr(shutil, "which", lambda name: "/usr/bin/soffice")
    calls = []
    monkeypatch.setattr(subprocess, "run", lambda args, **kw: calls.append(args))
    with pytest.raises(RuntimeError, match="no docx output"):
        office._soffice(tmp_path / "a.doc", "docx", tmp_path)
    (tmp_path / "a.docx").write_bytes(b"x")
    assert office._soffice(tmp_path / "a.doc", "docx", tmp_path) == tmp_path / "a.docx"
    assert calls[0][0] == "/usr/bin/soffice"


def test_engines_and_page_images(tmp_path: Path, make_docx, make_pdf):
    assert LOCAL[Category.READABLE_PDFS].label == "poppler@v2"
    assert Category.IMAGE_FILES not in LOCAL  # scans and photos are transcribed, not extracted
    docx = make_docx(tmp_path / "a.docx", "Pönnukökur")
    assert LOCAL[Category.DOCX_FILES].run(docx) == ("Pönnukökur\n", {})
    scan = make_pdf(tmp_path / "scan.pdf", ["x", "y"])
    pages = render_pages(scan, tmp_path / "pages", dpi=30)
    assert [p.name for p in pages] == ["scan-1.jpg", "scan-2.jpg"]
