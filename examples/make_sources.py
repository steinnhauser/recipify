"""Generate the sample documents in examples/sources/ from the markdown in examples/recipes/.

The documents imitate a real kitchen archive: Word files with soft line breaks and tables,
PDFs exported from Word (one in two columns), a photo of a recipe, a long handbook, an exact
copy, a .docx with its PDF twin, a mise en place list and a costing spreadsheet. The demo and
the end-to-end test run the pipeline on them.

Needs LibreOffice (soffice) and poppler (pdftoppm). Run from the repository root:

    uv run python examples/make_sources.py
"""

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Pt
from openpyxl import Workbook

ROOT = Path(__file__).parent
RECIPES, OUT = ROOT / "recipes", ROOT / "sources"


@dataclass
class Part:
    name: str = ""
    ingredients_label: str = ""
    ingredients: list[str] = field(default_factory=list)
    method_label: str = ""
    steps: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class Recipe:
    title: str
    intro: list[str] = field(default_factory=list)
    parts: list[Part] = field(default_factory=list)


def parse(path: Path) -> Recipe:
    """The small markdown dialect of examples/recipes (see examples/README.md)."""
    lines = path.read_text(encoding="utf-8").splitlines()
    recipe = Recipe(title=lines[0].removeprefix("# ").strip())
    part = Part()
    for line in (raw.strip() for raw in lines[1:]):
        if not line or re.fullmatch(r"\|[\s|:-]+\|", line) or line.startswith("| Amount"):
            continue
        if line.startswith("## ") and line[3:] not in ("Ingredients", "Method"):
            if part.ingredients or part.steps or part.name:
                recipe.parts.append(part)
            part = Part(name=line[3:])
        elif line.startswith("#"):
            label = line.lstrip("# ")
            if label.rstrip(":") in ("Ingredients", "Hráefni"):
                part.ingredients_label = label
            else:
                part.method_label = label
        elif line.startswith("- "):
            part.ingredients.append(line[2:])
        elif line.startswith("|"):
            part.ingredients.append(" ".join(c.strip() for c in line.strip("|").split("|")))
        elif line.startswith("> "):
            part.notes.append(line[2:])
        elif part.method_label or re.match(r"\d+\. ", line):
            part.steps.append(line)
        else:
            recipe.intro.append(line)
    recipe.parts.append(part)
    return recipe


def _two_columns(doc: Document) -> None:
    section = doc.sections[0]._sectPr
    cols = section.find(qn("w:cols"))
    if cols is None:
        cols = section.makeelement(qn("w:cols"), {})
        section.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "540")


def add_recipe(doc: Document, r: Recipe, headings: bool = False, table: bool = False) -> None:
    doc.add_heading(r.title, level=1)
    for line in r.intro:
        doc.add_paragraph(line)
    for part in r.parts:
        if part.name:
            if headings:
                doc.add_heading(part.name, level=2)
            else:
                doc.add_paragraph().add_run(part.name).bold = True
        if part.ingredients_label:
            doc.add_paragraph(part.ingredients_label)
        if table and part.ingredients:
            grid = doc.add_table(rows=0, cols=2)
            for line in part.ingredients:
                amount, _, name = line.partition(" ")
                if not re.match(r"[\d½]", amount):
                    amount, name = "", line
                cells = grid.add_row().cells
                cells[0].text, cells[1].text = amount, name
        elif part.ingredients:  # one paragraph, soft line breaks: how most kitchen files look
            run = doc.add_paragraph().add_run()
            for i, line in enumerate(part.ingredients):
                if i:
                    run.add_break()
                run.add_text(line)
        doc.add_paragraph(part.method_label or "Method")
        for step in part.steps:
            doc.add_paragraph(step)
        for note in part.notes:
            doc.add_paragraph(note)


def write_docx(path: Path, recipes: list[Recipe], columns: int = 1, **style: bool) -> Path:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(11)
    if columns == 2:
        _two_columns(doc)
    for i, r in enumerate(recipes):
        if i:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        add_recipe(doc, r, **style)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def to_pdf(docx: Path, target: Path) -> Path:
    with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            ["soffice", f"-env:UserInstallation=file://{profile}", "--headless",
             "--convert-to", "pdf", "--outdir", tmp, str(docx)],
            check=True, capture_output=True,
        )  # fmt: skip
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(Path(tmp) / f"{docx.stem}.pdf", target)
    return target


def pdf_from(recipes: list[Recipe], target: Path, columns: int = 1, **style: bool) -> Path:
    with tempfile.TemporaryDirectory() as tmp:
        return to_pdf(
            write_docx(Path(tmp) / f"{target.stem}.docx", recipes, columns, **style), target
        )


def to_png(pdf: Path, target: Path) -> None:
    stem = target.with_suffix("")
    subprocess.run(["pdftoppm", "-r", "70", "-png", "-singlefile", str(pdf), str(stem)], check=True)


def main() -> None:
    shutil.rmtree(OUT, ignore_errors=True)
    md = {p.stem: parse(p) for p in sorted(RECIPES.rglob("*.md"))}
    write_docx(OUT / "Bakery/Kleinur.docx", [md["kleinur"]])
    pdf_from([md["ponnukokur"]], OUT / "Breakfast/Pönnukökur.pdf")
    write_docx(OUT / "Drinks/Rhubarb cordial.docx", [md["rhubarb-cordial"]])
    write_docx(OUT / "Restaurant/Sauces/Béarnaise.docx", [md["bearnaise"]])
    write_docx(
        OUT / "Restaurant/Desserts/Skyr cheesecake.docx", [md["skyr-cheesecake"]], table=True
    )
    mousse = write_docx(OUT / "Restaurant/Desserts/Chocolate mousse.docx", [md["chocolate-mousse"]])
    to_pdf(mousse, OUT / "Restaurant/Desserts/Chocolate mousse.pdf")  # a .docx and its PDF twin
    pdf_from([md["skyr-og-blaber"]], OUT / "Restaurant/Eftirréttir/Skyr og bláber.pdf", columns=2)
    write_docx(
        OUT / "Restaurant/Mains/Cod with brown butter.docx", [md["cod-brown-butter"]], headings=True
    )
    for stem, name in [("lamb-shank", "Lamb shank"), ("mushroom-risotto", "Mushroom risotto")]:
        pdf_from([md[stem]], OUT / f"Restaurant/Mains/{name}.pdf")
    plokkfiskur = pdf_from([md["plokkfiskur"]], OUT / "Restaurant/Mains/Plokkfiskur.pdf")
    (OUT / "Old").mkdir()
    shutil.copy(plokkfiskur, OUT / "Old/Plokkfiskur (1).pdf")  # an exact copy
    pdf_from([md["potato-gratin"]], OUT / "Restaurant/Sides/Potato gratin.pdf")
    pdf_from([md["smoked-trout-canape"]], OUT / "Restaurant/Canapés/Smoked trout canapé.pdf")
    with tempfile.TemporaryDirectory() as tmp:  # a photo of a printed recipe: needs OCR
        (OUT / "Scans").mkdir()
        to_png(
            pdf_from([md["beetroot-salad"]], Path(tmp) / "b.pdf"), OUT / "Scans/Beetroot salad.png"
        )
    pdf_from(list(md.values()), OUT / "Kitchen handbook 2019.pdf")  # a book: one page per recipe
    doc = Document()
    doc.add_heading("Missi listi – morgunvakt", level=1)
    for line in ["Skera lauk og gulrætur", "Þeyta rjóma", "Fylla á krydd", "Merkja og dagsetja"]:
        doc.add_paragraph(line)
    doc.save(str(OUT / "Missi listi.docx"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Costing"
    ws.append(["Item", "Quantity", "Unit", "Price/unit", "Cost"])
    for row in [("Lamb shank", 4, "pcs", 950, 3800), ("Red wine", 0.3, "l", 2400, 720)]:
        ws.append(list(row))
    wb.save(str(OUT / "Costing - lamb shank.xlsx"))
    for path in sorted(OUT.rglob("*")):
        if path.is_file():
            print(path.relative_to(ROOT))


if __name__ == "__main__":
    main()
