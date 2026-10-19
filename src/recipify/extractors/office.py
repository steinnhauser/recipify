"""Spreadsheets and legacy office documents, converted locally (LibreOffice for old formats)."""

import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from recipify.extractors.docx import docx_to_markdown
from recipify.extractors.pdf import pdf_to_markdown

MAX_ROWS = 2000


def _soffice(path: Path, target: str, outdir: Path) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) is required to convert this format")
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [
                exe,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                target,
                "--outdir",
                str(outdir),
                str(path),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
    out = outdir / f"{path.stem}.{target}"
    if not out.exists():
        raise RuntimeError(f"LibreOffice produced no {target} output")
    return out


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    elif isinstance(value, float):
        value = round(value, 4)
    return str(value).strip().replace("\n", " ").replace("|", "\\|")


def _xlsx(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    blocks = [f"# {path.stem}"]
    for ws in wb.worksheets:
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS)]
        rows = [r for r in rows if any(r)]
        if not rows:
            continue
        width = max(i for r in rows for i, c in enumerate(r, 1) if c)
        rows = [(r + [""] * width)[:width] for r in rows]
        lines = ["| " + " | ".join(rows[0]) + " |", "|" + " --- |" * width]
        lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        blocks.append(f"## {ws.title}\n\n" + "\n".join(lines))
    wb.close()
    return "\n\n".join(blocks) + "\n"


def spreadsheet_to_markdown(path: Path) -> str:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _xlsx(path)
    with tempfile.TemporaryDirectory() as tmp:
        return _xlsx(_soffice(path, "xlsx", Path(tmp)))


def legacy_office_to_markdown(path: Path) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        if path.suffix.lower() in {".ppt", ".pptx", ".odp"}:
            return pdf_to_markdown(_soffice(path, "pdf", Path(tmp)))
        text = docx_to_markdown(_soffice(path, "docx", Path(tmp)))
    return f"# {path.stem}\n\n{text}" if not text.startswith("#") else text
