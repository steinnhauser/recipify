"""Decide each file's category from its content, not just its name."""

import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from recipify.categories import Category
from recipify.extract.pdf import page_count, pdf_text

MIN_PDF_TEXT_CHARS = 100  # below this over the sampled pages, the PDF is a scan
PDF_SAMPLE_PAGES = 5

IMAGE_MAGIC = (b"\xff\xd8\xff", b"\x89PNG\r\n\x1a\n")
IGNORED_NAMES = {".ds_store", "thumbs.db", "desktop.ini"}
IGNORED_SUFFIXES = {".url", ".zip", ".lnk", ".tmp"}
IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png"}
SPREADSHEET_SUFFIXES = {".xls", ".xlsx", ".csv"}
LEGACY_SUFFIXES = {".doc", ".odt", ".pptx", ".ppt", ".rtf"}


@dataclass(frozen=True)
class Classification:
    category: Category
    reason: str
    pages: int | None = None


def _head(path: Path, n: int = 8) -> bytes:
    with path.open("rb") as f:
        return f.read(n)


def _sniff_zip(path: Path) -> Classification | None:
    if not zipfile.is_zipfile(path):
        return None
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
    if "word/document.xml" in names:
        return _classify_docx(path)
    if "META-INF/container.xml" in names or "mimetype" in names:
        return Classification(Category.EPUB_BOOKS, "EPUB container (misnamed)")
    return None


def _classify_pdf(path: Path) -> Classification:
    if not _head(path, 5).startswith(b"%PDF"):
        sniffed = _sniff_zip(path)  # misnamed files: an EPUB or docx called .pdf
        return sniffed or Classification(Category.EMPTY_OR_CORRUPTED, "missing %PDF header")
    try:
        pages = page_count(path)
        text = "".join(pdf_text(path, "raw", PDF_SAMPLE_PAGES).split())
    except RuntimeError as exc:  # encrypted, truncated or otherwise unreadable
        return Classification(Category.EMPTY_OR_CORRUPTED, f"unreadable PDF: {exc}")
    if pages == 0:
        return Classification(Category.EMPTY_OR_CORRUPTED, "PDF has no pages")
    if len(text) >= MIN_PDF_TEXT_CHARS:
        return Classification(Category.READABLE_PDFS, f"{len(text)} chars of text", pages)
    return Classification(Category.IMAGE_BASED_PDFS, f"only {len(text)} chars of text", pages)


def _classify_docx(path: Path) -> Classification:
    if not zipfile.is_zipfile(path):
        return Classification(Category.EMPTY_OR_CORRUPTED, "not a valid docx (zip) container")
    try:
        doc = Document(str(path))
    except Exception as exc:  # python-docx raises many types on damaged files
        return Classification(Category.EMPTY_OR_CORRUPTED, f"unreadable docx: {exc}")
    has_text = any(p.text.strip() for p in doc.paragraphs) or any(
        c.text.strip() for t in doc.tables for r in t.rows for c in r.cells
    )
    if not has_text:
        return Classification(Category.EMPTY_OR_CORRUPTED, "docx contains no text")
    return Classification(Category.DOCX_FILES, "valid docx with text")


def classify(path: Path) -> Classification:
    name, suffix = path.name.lower(), path.suffix.lower()
    if name in IGNORED_NAMES or name.startswith(("~$", ".")) or suffix in IGNORED_SUFFIXES:
        return Classification(Category.IGNORED, "lock/system/shortcut/archive file")
    if path.stat().st_size == 0:
        return Classification(Category.EMPTY_OR_CORRUPTED, "zero bytes")
    if suffix == ".epub":
        return Classification(Category.EPUB_BOOKS, "EPUB container")
    if suffix == ".pdf":
        return _classify_pdf(path)
    if suffix == ".docx":
        return _classify_docx(path)
    if suffix in IMAGE_SUFFIXES:
        if _head(path).startswith(IMAGE_MAGIC):
            return Classification(Category.IMAGE_FILES, "valid image header")
        return Classification(Category.EMPTY_OR_CORRUPTED, "image header not recognised")
    if suffix in SPREADSHEET_SUFFIXES:
        return Classification(Category.SPREADSHEETS, f"{suffix} spreadsheet")
    if suffix in LEGACY_SUFFIXES:
        return Classification(Category.LEGACY_OFFICE, f"{suffix} needs conversion")
    return Classification(Category.OTHER_FORMATS, f"unrecognised extension {suffix or '(none)'}")
