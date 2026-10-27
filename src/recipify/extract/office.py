"""Spreadsheets and legacy office documents, converted locally (LibreOffice for old formats)."""

import re
import shutil
import subprocess
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from recipify.extract.docx import extract_docx
from recipify.extract.pdf import extract_pdf

SPREADSHEET_ENGINE, SPREADSHEET_VERSION = "openpyxl", 2
LEGACY_ENGINE, LEGACY_VERSION = "libreoffice", 2
MAX_ROWS = 2000
# Sheet names that say nothing about the content: "1", "Sheet5", "Blað2" (Icelandic Excel).
GENERIC_SHEET = re.compile(r"^(?:\d+|sheet\s*\d*|blad\s*\d*|tafla\s*\d*|feuil\w*\s*\d*)$", re.I)


def _soffice(path: Path, target: str, outdir: Path) -> Path:
    exe = shutil.which("soffice") or shutil.which("libreoffice")
    if not exe:
        raise RuntimeError("LibreOffice (soffice) is required to convert this format")
    with tempfile.TemporaryDirectory() as profile:
        subprocess.run(
            [
                exe,
                f"-env:UserInstallation=file://{profile}",
                "--headless",
                "--convert-to",
                target,
                "--outdir",
                str(outdir),
                str(path),
            ],
            check=True,
            capture_output=True,
            timeout=180,
        )
    out = outdir / f"{path.stem}.{target}"
    if not out.exists():
        raise RuntimeError(f"LibreOffice produced no {target} output")
    return out


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    elif isinstance(value, float):
        value = round(value, 4)
    return str(value).strip().replace("\n", " ").replace("|", "/")


def _column_blocks(rows: list[list[str]]) -> list[list[int]]:
    """Groups of neighbouring columns, separated by at least one empty column."""
    used = sorted({i for row in rows for i, cell in enumerate(row) if cell})
    blocks: list[list[int]] = []
    for col in used:
        if blocks and col == blocks[-1][-1] + 1:
            blocks[-1].append(col)
        else:
            blocks.append([col])
    return blocks


def _sheet(rows: list[list[str]]) -> list[str]:
    """Each column block read top to bottom; within a block, a row reads left to right.

    Two recipes laid out side by side (columns B and E) stay apart, while a row such as
    "500 g | Milk | 500 g | Mjólk" keeps its quantity, name and translation together.
    """
    lines = []
    for block in _column_blocks(rows):
        cells = [[row[i] for i in block if i < len(row) and row[i]] for row in rows]
        lines += [" | ".join(c) for c in cells if c]
    return lines


def _xlsx(path: Path, title: str) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    pages = []
    for ws in wb.worksheets:
        rows = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS)]
        lines = _sheet(rows)
        if lines:
            name = ws.title.strip()
            meaningful = not GENERIC_SHEET.match(name.replace("ð", "d"))
            pages.append((f"## {name}\n\n" if meaningful else "") + "\n".join(lines))
    wb.close()
    return f"# {title}\n\n" + "\n\f\n".join(pages) + "\n"  # one page per sheet


def extract_spreadsheet(path: Path) -> str:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return _xlsx(path, path.stem)
    with tempfile.TemporaryDirectory() as tmp:
        return _xlsx(_soffice(path, "xlsx", Path(tmp)), path.stem)


def extract_legacy(path: Path) -> str:
    with tempfile.TemporaryDirectory() as tmp:
        if path.suffix.lower() in {".ppt", ".pptx", ".odp"}:
            return extract_pdf(_soffice(path, "pdf", Path(tmp)))[0]
        return extract_docx(_soffice(path, "docx", Path(tmp)))
